from __future__ import annotations

import csv
import json
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


REPORT_DIR = (
    Path(__file__).resolve().parents[2]
    / "backend"
    / "generated_reports"
)

REPORT_DIR.mkdir(parents=True, exist_ok=True)


class ReportAgent:
    """
    Multi-format report generation engine.

    Supported formats:
    - PDF
    - Excel
    - CSV
    - JSON
    """

    SUPPORTED_FORMATS = {
        "pdf",
        "excel",
        "xlsx",
        "csv",
        "json",
    }

    def __init__(self, output_dir: Path | None = None):
        self.output_dir = output_dir or REPORT_DIR
        self.output_dir.mkdir(
            parents=True,
            exist_ok=True,
        )

    # ==========================================================
    # NORMALIZE DATA
    # ==========================================================

    def normalize_payload(
        self,
        payload: dict[str, Any],
    ) -> dict[str, Any]:

        return {
            "report_title": payload.get(
                "report_title",
                "TechNova Sentinel AI Health Report",
            ),
            "report_type": payload.get(
                "report_type",
                "on-demand",
            ),
            "generated_at": payload.get(
                "generated_at",
                datetime.now().isoformat(),
            ),
            "district": payload.get(
                "district",
                "Tamil Nadu",
            ),
            "disease": payload.get(
                "disease",
                "Dengue",
            ),
            "prediction_summary": payload.get(
                "prediction_summary",
                {},
            ),
            "weather_trends": payload.get(
                "weather_trends",
                [],
            ),
            "hospital_resource_status": payload.get(
                "hospital_resource_status",
                [],
            ),
            "recommendations": payload.get(
                "recommendations",
                [],
            ),
        }

    # ==========================================================
    # FILE NAME
    # ==========================================================

    def build_filename(
        self,
        report_type: str,
        file_format: str,
    ) -> str:

        timestamp = datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )

        extension = (
            "xlsx"
            if file_format == "excel"
            else file_format
        )

        return (
            f"technova_{report_type}_"
            f"{timestamp}.{extension}"
        )

    # ==========================================================
    # JSON
    # ==========================================================

    def generate_json(
        self,
        data: dict[str, Any],
    ) -> Path:

        filename = self.build_filename(
            data["report_type"],
            "json",
        )

        path = self.output_dir / filename

        with path.open(
            "w",
            encoding="utf-8",
        ) as file:

            json.dump(
                data,
                file,
                indent=2,
                ensure_ascii=False,
                default=str,
            )

        return path

    # ==========================================================
    # CSV
    # ==========================================================

    def generate_csv(
        self,
        data: dict[str, Any],
    ) -> Path:

        filename = self.build_filename(
            data["report_type"],
            "csv",
        )

        path = self.output_dir / filename

        rows = []

        prediction = data.get(
            "prediction_summary",
            {},
        )

        rows.append(
            {
                "section": "prediction_summary",
                "metric": "district",
                "value": data.get("district"),
            }
        )

        rows.append(
            {
                "section": "prediction_summary",
                "metric": "disease",
                "value": data.get("disease"),
            }
        )

        for key, value in prediction.items():

            rows.append(
                {
                    "section": "prediction_summary",
                    "metric": key,
                    "value": value,
                }
            )

        for item in data.get(
            "weather_trends",
            [],
        ):

            if isinstance(item, dict):

                for key, value in item.items():

                    rows.append(
                        {
                            "section": "weather_trends",
                            "metric": key,
                            "value": value,
                        }
                    )

        for item in data.get(
            "hospital_resource_status",
            [],
        ):

            if isinstance(item, dict):

                for key, value in item.items():

                    rows.append(
                        {
                            "section": "hospital_resource_status",
                            "metric": key,
                            "value": value,
                        }
                    )

        for index, recommendation in enumerate(
            data.get("recommendations", []),
            start=1,
        ):

            rows.append(
                {
                    "section": "recommendations",
                    "metric": f"recommendation_{index}",
                    "value": recommendation,
                }
            )

        dataframe = pd.DataFrame(rows)

        dataframe.to_csv(
            path,
            index=False,
            encoding="utf-8-sig",
        )

        return path

    # ==========================================================
    # EXCEL
    # ==========================================================

    def generate_excel(
        self,
        data: dict[str, Any],
    ) -> Path:

        filename = self.build_filename(
            data["report_type"],
            "excel",
        )

        path = self.output_dir / filename

        workbook = Workbook()

        summary_sheet = workbook.active
        summary_sheet.title = "Prediction Summary"

        summary_sheet.append(
            [
                "Metric",
                "Value",
            ]
        )

        summary_sheet.append(
            [
                "District",
                data.get("district"),
            ]
        )

        summary_sheet.append(
            [
                "Disease",
                data.get("disease"),
            ]
        )

        for key, value in data.get(
            "prediction_summary",
            {},
        ).items():

            summary_sheet.append(
                [
                    key,
                    value,
                ]
            )

        # ------------------------------------------------------

        weather_sheet = workbook.create_sheet(
            "Weather Trends"
        )

        weather_rows = data.get(
            "weather_trends",
            [],
        )

        if weather_rows:

            headers = list(
                weather_rows[0].keys()
            )

            weather_sheet.append(headers)

            for row in weather_rows:

                weather_sheet.append(
                    [
                        row.get(header)
                        for header in headers
                    ]
                )

        else:

            weather_sheet.append(
                [
                    "No weather trend data"
                ]
            )

        # ------------------------------------------------------

        hospital_sheet = workbook.create_sheet(
            "Hospital Resources"
        )

        hospital_rows = data.get(
            "hospital_resource_status",
            [],
        )

        if hospital_rows:

            headers = list(
                hospital_rows[0].keys()
            )

            hospital_sheet.append(headers)

            for row in hospital_rows:

                hospital_sheet.append(
                    [
                        row.get(header)
                        for header in headers
                    ]
                )

        else:

            hospital_sheet.append(
                [
                    "No hospital resource data"
                ]
            )

        # ------------------------------------------------------

        recommendation_sheet = workbook.create_sheet(
            "Recommendations"
        )

        recommendation_sheet.append(
            [
                "No.",
                "Recommendation",
            ]
        )

        for index, recommendation in enumerate(
            data.get("recommendations", []),
            start=1,
        ):

            recommendation_sheet.append(
                [
                    index,
                    recommendation,
                ]
            )

        # ------------------------------------------------------

        for worksheet in workbook.worksheets:

            worksheet.freeze_panes = "A2"

            for column in worksheet.columns:

                max_length = 0

                column_letter = column[0].column_letter

                for cell in column:

                    value = (
                        ""
                        if cell.value is None
                        else str(cell.value)
                    )

                    max_length = max(
                        max_length,
                        len(value),
                    )

                worksheet.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 2,
                    50,
                )

        workbook.save(path)

        return path

    # ==========================================================
    # PDF
    # ==========================================================

    def generate_pdf(
        self,
        data: dict[str, Any],
    ) -> Path:

        filename = self.build_filename(
            data["report_type"],
            "pdf",
        )

        path = self.output_dir / filename

        document = SimpleDocTemplate(
            str(path),
            pagesize=A4,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36,
        )

        styles = getSampleStyleSheet()

        story = []

        story.append(
            Paragraph(
                data["report_title"],
                styles["Title"],
            )
        )

        story.append(
            Spacer(1, 12)
        )

        story.append(
            Paragraph(
                f"<b>District:</b> "
                f"{data.get('district')}",
                styles["Normal"],
            )
        )

        story.append(
            Paragraph(
                f"<b>Disease:</b> "
                f"{data.get('disease')}",
                styles["Normal"],
            )
        )

        story.append(
            Paragraph(
                f"<b>Report Type:</b> "
                f"{data.get('report_type')}",
                styles["Normal"],
            )
        )

        story.append(
            Paragraph(
                f"<b>Generated At:</b> "
                f"{data.get('generated_at')}",
                styles["Normal"],
            )
        )

        story.append(
            Spacer(1, 18)
        )

        # Prediction Summary

        story.append(
            Paragraph(
                "Prediction Summary",
                styles["Heading2"],
            )
        )

        prediction_rows = [
            [
                "Metric",
                "Value",
            ]
        ]

        for key, value in data.get(
            "prediction_summary",
            {},
        ).items():

            prediction_rows.append(
                [
                    str(key),
                    str(value),
                ]
            )

        if len(prediction_rows) == 1:

            prediction_rows.append(
                [
                    "Status",
                    "No prediction data",
                ]
            )

        prediction_table = Table(
            prediction_rows,
            repeatRows=1,
        )

        prediction_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor("#1e293b"),
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white,
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.grey,
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                ]
            )
        )

        story.append(prediction_table)

        story.append(
            Spacer(1, 18)
        )

        # Weather

        story.append(
            Paragraph(
                "Weather Trends",
                styles["Heading2"],
            )
        )

        weather = data.get(
            "weather_trends",
            [],
        )

        if weather:

            weather_rows = [
                list(weather[0].keys())
            ]

            for row in weather:

                weather_rows.append(
                    [
                        str(value)
                        for value in row.values()
                    ]
                )

            weather_table = Table(
                weather_rows,
                repeatRows=1,
            )

            weather_table.setStyle(
                TableStyle(
                    [
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.HexColor("#1e293b"),
                        ),
                        (
                            "TEXTCOLOR",
                            (0, 0),
                            (-1, 0),
                            colors.white,
                        ),
                        (
                            "GRID",
                            (0, 0),
                            (-1, -1),
                            0.5,
                            colors.grey,
                        ),
                    ]
                )
            )

            story.append(weather_table)

        else:

            story.append(
                Paragraph(
                    "No weather trend data available.",
                    styles["Normal"],
                )
            )

        story.append(
            Spacer(1, 18)
        )

        # Hospital Resources

        story.append(
            Paragraph(
                "Hospital Resource Status",
                styles["Heading2"],
            )
        )

        hospital_rows = data.get(
            "hospital_resource_status",
            [],
        )

        if hospital_rows:

            rows = [
                list(hospital_rows[0].keys())
            ]

            for row in hospital_rows:

                rows.append(
                    [
                        str(value)
                        for value in row.values()
                    ]
                )

            hospital_table = Table(
                rows,
                repeatRows=1,
            )

            hospital_table.setStyle(
                TableStyle(
                    [
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.HexColor("#1e293b"),
                        ),
                        (
                            "TEXTCOLOR",
                            (0, 0),
                            (-1, 0),
                            colors.white,
                        ),
                        (
                            "GRID",
                            (0, 0),
                            (-1, -1),
                            0.5,
                            colors.grey,
                        ),
                    ]
                )
            )

            story.append(
                hospital_table
            )

        else:

            story.append(
                Paragraph(
                    "No hospital resource data available.",
                    styles["Normal"],
                )
            )

        story.append(
            Spacer(1, 18)
        )

        # Recommendations

        story.append(
            Paragraph(
                "Recommendations",
                styles["Heading2"],
            )
        )

        recommendations = data.get(
            "recommendations",
            [],
        )

        if recommendations:

            for index, recommendation in enumerate(
                recommendations,
                start=1,
            ):

                story.append(
                    Paragraph(
                        f"{index}. "
                        f"{recommendation}",
                        styles["Normal"],
                    )
                )

        else:

            story.append(
                Paragraph(
                    "No recommendations available.",
                    styles["Normal"],
                )
            )

        document.build(story)

        return path

    # ==========================================================
    # GENERATE
    # ==========================================================

    def generate(
        self,
        payload: dict[str, Any],
        file_format: str,
    ) -> dict[str, Any]:

        normalized = self.normalize_payload(
            payload
        )

        requested_format = (
            str(file_format)
            .lower()
            .strip()
        )

        if requested_format not in self.SUPPORTED_FORMATS:

            raise ValueError(
                "Unsupported report format. "
                "Use PDF, Excel, CSV, or JSON."
            )

        if requested_format == "pdf":

            path = self.generate_pdf(
                normalized
            )

        elif requested_format in {
            "excel",
            "xlsx",
        }:

            path = self.generate_excel(
                normalized
            )

        elif requested_format == "csv":

            path = self.generate_csv(
                normalized
            )

        elif requested_format == "json":

            path = self.generate_json(
                normalized
            )

        else:

            raise ValueError(
                "Unsupported report format."
            )

        return {
            "filename": path.name,
            "path": str(path),
            "format": (
                "excel"
                if requested_format == "xlsx"
                else requested_format
            ),
            "report_type": normalized[
                "report_type"
            ],
            "generated_at": normalized[
                "generated_at"
            ],
        }